#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel(.xlsx/.xlsm) 取り込み — 自治会会計管理システム（パソコン版）用

スマホ版「自治会会計」が書き出した Excel や、このアプリ自身が出した Excel
（概要・残高／取引一覧／振替／科目別集計）を読み込んで、kaikei.db に入れる。

・見出しの行は自分でさがす（表題が上に何行あってもよい）
・見出しの言い方が違っても当てはめる（年月日／勘定科目／収入金額／支払方法 …）
・日付は 2026-04-05・2026/4/5・R8.6.1・令和8年6月1日・Excelの日付 のどれでも読む
・金額は 8,800・¥8,800・▲5,000・(5,000)・全角 も読む
・すでに入っている取引は二重に入れない。ただし、まったく同じ内容の取引が
  2件ある帳面もあるので、「何件あるか」で数えて突き合わせる（1件に丸めない）
・新しい科目・口座は自動で足す
・途中で失敗したら、ぜんぶ取り消す（中途半端に入らない）

main.py からの使い方（2か所だけ）:

    from excel_import import import_excel_dialog        # ファイルの先頭あたり

    # 「データ ▼」メニューに1行足す（CSV取り込みの下あたり）
    menu.add_command(label='Excelを取り込み', command=lambda: import_excel_dialog(self))

app には .root / .db / .current_fy（IntVar）があればよい。
取り込みのあと、あれば refresh_list・refresh_summary・refresh_balance を呼ぶ。
"""

import re
import unicodedata
from datetime import datetime, date

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:                                   # openpyxl が無い環境でも import は通す
    EXCEL_AVAILABLE = False


# ============================================================
# 見出しの言い方
# ============================================================
ALIASES = {
    'date':  ['日付', '年月日', '取引日', '取引年月日', '月日', '日時', '起票日', '日'],
    'acc':   ['口座', '口座名', '出どころ', '支払方法', '支払区分', '現金口座', '現金・口座',
              '預金', '金融機関', '取引口座'],
    'kind':  ['区分', '収支', '収支区分', '種別', '入出金', '入出金区分', '収入支出'],
    'cat':   ['科目', '勘定科目', '費目', '項目', '細目', '分類', '種目'],
    'note':  ['摘要', '内容', '取引内容', '明細', '品目', '相手先', '支払先', '用途', '件名'],
    'income': ['収入', '収入金額', '収入額', '入金', '入金額', '受入', '受入金額',
               '借方', '借方金額', '入'],
    'expense': ['支出', '支出金額', '支出額', '出金', '出金額', '払出', '払出金額',
                '貸方', '貸方金額', '出'],
    'amount': ['金額', '額', '合計金額', '取引金額'],
    'balance': ['残高', '差引残高', '差引', '繰越残高', '現在残高', '残'],
    'event': ['行事', '行事名', '事業', '事業名', 'イベント'],
    'memo':  ['備考', 'メモ', '注記', 'コメント'],
    'from':  ['振替元', '振替元口座', '出金元', '引出元', '移動元'],
    'to':    ['振替先', '振替先口座', '入金先', '預入先', '移動先'],
    'tramt': ['振替額', '振替金額', '移動額'],
    'opening': ['期首残高', '期首', '前期繰越', '前年度繰越', '繰越金'],
}
TX_FIELDS = ('date', 'acc', 'kind', 'cat', 'note', 'income', 'expense',
             'amount', 'balance', 'event', 'memo')
TR_FIELDS = ('date', 'from', 'to', 'amount', 'tramt', 'note', 'memo')

WAREKI = {'令和': 2018, 'r': 2018, '平成': 1988, 'h': 1988,
          '昭和': 1925, 's': 1925, '大正': 1911, 't': 1911, '明治': 1867, 'm': 1867}


def _norm(s):
    """見出しをくらべるための形にそろえる（全角→半角・空き・かっこを落とす）"""
    if s is None:
        return ''
    s = unicodedata.normalize('NFKC', str(s))
    s = re.sub(r'[（(].*?[)）]', '', s)
    s = re.sub(r'[\s　()\[\]【】]', '', s)
    return s.strip().lower()


def _field_of(header, allow):
    n = _norm(header)
    if not n:
        return None
    for f in allow:                                   # まず、ぴったり同じ言い方
        if any(_norm(a) == n for a in ALIASES.get(f, [])):
            return f
    for f in allow:                                   # つぎに、含んでいるもの
        if any(len(_norm(a)) >= 2 and _norm(a) in n for a in ALIASES.get(f, [])):
            return f
    return None


# ============================================================
# 値の読み取り
# ============================================================
def parse_date(v, fiscal_year=None):
    """いろいろな書き方の日付を 'YYYY-MM-DD' にする。読めなければ ''"""
    if v is None or v == '':
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        n = float(v)
        if 20000 <= n <= 80000:                       # Excel の日付（1899-12-30 からの日数）
            from datetime import timedelta
            return (datetime(1899, 12, 30) + timedelta(days=int(n))).strftime('%Y-%m-%d')
        if 19000101 <= n <= 21001231:                 # 20260405 のような8けた
            s = str(int(n))
            return f'{s[0:4]}-{s[4:6]}-{s[6:8]}'
        return ''
    s = unicodedata.normalize('NFKC', str(v)).strip()
    m = re.match(r'^(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})', s)
    if m:
        return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'
    m = re.match(r'^(令和|平成|昭和|大正|明治|[RrHhSsTtMm])\s*(\d{1,2}|元)[-/年.](\d{1,2})[-/月.](\d{1,2})', s)
    if m:
        base = WAREKI.get(m.group(1).lower())
        if base:
            yy = 1 if m.group(2) == '元' else int(m.group(2))
            return f'{base + yy:04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}'
    m = re.match(r'^(\d{1,2})[-/月.](\d{1,2})', s)     # 年のない「4/5」は年度から補う
    if m and fiscal_year:
        mo, da = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12 and 1 <= da <= 31:
            y = fiscal_year if mo >= 4 else fiscal_year + 1
            return f'{y:04d}-{mo:02d}-{da:02d}'
    return ''


def parse_money(v):
    """8,800 / ¥8,800 / ▲5,000 / (5,000) / 全角 を int にする"""
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return int(round(float(v)))
    s = unicodedata.normalize('NFKC', str(v))
    s = re.sub(r'[,\s　¥円]', '', s)
    sign = 1
    if re.match(r'^[(（].*[)）]$', s):
        sign, s = -1, s[1:-1]
    if re.match(r'^[▲△−–—-]', s):
        sign, s = -1, re.sub(r'^[▲△−–—-]', '', s)
    try:
        return int(round(float(s))) * sign
    except ValueError:
        return 0


def parse_kind(v):
    s = _norm(v)
    if not s:
        return None
    if re.search(r'収入|入金|受入|借方|収|入|\+', s):
        return 'in'
    if re.search(r'支出|出金|払出|貸方|支|出|-|▲|△', s):
        return 'out'
    return None


def fiscal_year_of(date_str):
    """年度は4月はじまり（1〜3月は前の年の年度）"""
    m = re.match(r'^(\d{4})-(\d{2})', date_str or '')
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    return y if mo >= 4 else y - 1


# ============================================================
# シートを読む
# ============================================================
def _grid(ws, max_rows=20000, max_cols=64):
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows),
                          max_col=min(ws.max_column, max_cols), values_only=True):
        rows.append(list(r))
    return rows


def _detect_header(grid, allow, need):
    """見出しの行と、どの列が何かを当てる。戻り: (行番号, {項目: 列番号}, 手ごたえ)"""
    best = (-1, {}, 0)
    for r in range(min(len(grid), 30)):
        row = grid[r]
        mapping, used, score = {}, set(), 0
        for c, cell in enumerate(row):
            if cell is None or cell == '' or isinstance(cell, (int, float)):
                continue
            f = _field_of(cell, allow)
            if f and f not in mapping and c not in used:
                mapping[f] = c
                used.add(c)
                score += 1
        s = score + (10 if need(mapping) else 0)
        if s > best[2]:
            best = (r, mapping, s)
    return best


def _need_tx(m):
    return 'date' in m and ('income' in m or 'expense' in m or 'amount' in m)


def _need_tr(m):
    return 'date' in m and 'from' in m and 'to' in m and ('tramt' in m or 'amount' in m)


def _cell(row, mapping, field):
    c = mapping.get(field)
    if c is None or c >= len(row):
        return None
    return row[c]


def _text(row, mapping, field):
    v = _cell(row, mapping, field)
    return '' if v is None else str(v).strip()


def read_transactions(grid, header_row, mapping, fiscal_year=None):
    """取引の表 → [{date, account, category, details, income, expense, event, memo}]"""
    out, skipped = [], []
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        if not any(c is not None and c != '' for c in row):
            continue
        d = parse_date(_cell(row, mapping, 'date'), fiscal_year)
        inc = parse_money(_cell(row, mapping, 'income')) if 'income' in mapping else 0
        exp = parse_money(_cell(row, mapping, 'expense')) if 'expense' in mapping else 0
        if inc <= 0 and exp <= 0 and 'amount' in mapping:
            a = parse_money(_cell(row, mapping, 'amount'))
            k = parse_kind(_cell(row, mapping, 'kind')) if 'kind' in mapping else None
            if a:
                if k == 'out' or (k is None and a < 0):
                    exp = abs(a)
                else:
                    inc = abs(a)
        if 'kind' in mapping and (inc > 0 or exp > 0):
            k = parse_kind(_cell(row, mapping, 'kind'))
            if k == 'out' and inc > 0:
                inc, exp = 0, inc
            elif k == 'in' and exp > 0:
                inc, exp = exp, 0
        if not d:
            skipped.append((r + 1, '日付が読めません'))
            continue
        if inc <= 0 and exp <= 0:
            skipped.append((r + 1, '金額がありません'))
            continue
        out.append({
            'date':     d,
            'account':  _text(row, mapping, 'acc'),
            'category': _text(row, mapping, 'cat'),
            'details':  _text(row, mapping, 'note'),
            'income':   inc if inc > 0 else 0,
            'expense':  exp if exp > 0 else 0,
            'event':    _text(row, mapping, 'event'),
            'memo':     _text(row, mapping, 'memo'),
            'row':      r + 1,
        })
    return out, skipped


def read_transfers(grid, header_row, mapping, fiscal_year=None):
    """振替の表 → [{date, from_account, to_account, amount, details, memo}]"""
    out, skipped = [], []
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        if not any(c is not None and c != '' for c in row):
            continue
        d = parse_date(_cell(row, mapping, 'date'), fiscal_year)
        amt = abs(parse_money(_cell(row, mapping, 'tramt') if 'tramt' in mapping
                              else _cell(row, mapping, 'amount')))
        fr, to = _text(row, mapping, 'from'), _text(row, mapping, 'to')
        if not d:
            skipped.append((r + 1, '日付が読めません'))
            continue
        if not amt or not fr or not to:
            skipped.append((r + 1, '振替元・振替先・金額がそろっていません'))
            continue
        out.append({'date': d, 'from_account': fr, 'to_account': to, 'amount': amt,
                    'details': _text(row, mapping, 'note'), 'memo': _text(row, mapping, 'memo'),
                    'row': r + 1})
    return out, skipped


def read_opening_balances(grid):
    """「口座／期首残高」の表 → {口座名: 金額}。見つからなければ {}"""
    for r in range(min(len(grid), 40)):
        row = grid[r]
        if len(row) < 2:
            continue
        if _field_of(row[0], ['acc']) == 'acc' and _field_of(row[1], ['opening']) == 'opening':
            out = {}
            for k in range(r + 1, len(grid)):
                rr = grid[k]
                name = '' if not rr or rr[0] is None else str(rr[0]).strip()
                if not name or _norm(name) == '合計':
                    break
                out[name] = parse_money(rr[1] if len(rr) > 1 else 0)
            return out
    return {}


# ============================================================
# ブック全体を読む
# ============================================================
class ImportPlan:
    """読み取った中身と、入れたときにどうなるかの見立て"""

    def __init__(self):
        self.transactions = []
        self.transfers = []
        self.openings = {}
        self.skipped = []
        self.fiscal_years = set()
        self.new_categories = []
        self.new_accounts = []
        self.dup_transactions = 0
        self.dup_transfers = 0

    @property
    def add_transactions(self):
        return len(self.transactions) - self.dup_transactions

    @property
    def add_transfers(self):
        return len(self.transfers) - self.dup_transfers


def read_workbook(path, fiscal_year=None):
    """.xlsx/.xlsm を読んで ImportPlan（重複の数はまだ入らない）を返す"""
    if not EXCEL_AVAILABLE:
        raise RuntimeError('openpyxl が入っていないため、Excel を読み込めません。')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    plan = ImportPlan()
    tx_best = (None, -1)
    tr_best = (None, -1)
    try:
        for ws in wb.worksheets:
            grid = _grid(ws)
            if not grid:
                continue
            hr, mp, sc = _detect_header(grid, TX_FIELDS, _need_tx)
            if _need_tx(mp) and sc > tx_best[1]:
                tx_best = ((grid, hr, mp), sc)
            hr2, mp2, sc2 = _detect_header(grid, TR_FIELDS, _need_tr)
            if _need_tr(mp2) and sc2 > tr_best[1]:
                tr_best = ((grid, hr2, mp2), sc2)
            if not plan.openings:
                plan.openings = read_opening_balances(grid)
    finally:
        wb.close()

    if tx_best[0]:
        grid, hr, mp = tx_best[0]
        plan.transactions, sk = read_transactions(grid, hr, mp, fiscal_year)
        plan.skipped += [('取引', r, w) for r, w in sk]
    if tr_best[0]:
        grid, hr, mp = tr_best[0]
        plan.transfers, sk = read_transfers(grid, hr, mp, fiscal_year)
        plan.skipped += [('振替', r, w) for r, w in sk]
    for t in plan.transactions:
        y = fiscal_year_of(t['date'])
        if y is not None:
            plan.fiscal_years.add(y)
    for t in plan.transfers:
        y = fiscal_year_of(t['date'])
        if y is not None:
            plan.fiscal_years.add(y)
    return plan


# ============================================================
# いまの中身とくらべる
# ============================================================
def _tx_key(t):
    return (t['date'], (t.get('account') or '').strip(), (t.get('category') or '').strip(),
            (t.get('details') or '').strip(), int(round(t.get('income') or 0)),
            int(round(t.get('expense') or 0)))


def _tr_key(t):
    return (t['date'], (t.get('from_account') or '').strip(),
            (t.get('to_account') or '').strip(), int(round(t.get('amount') or 0)))


def _db_tx_key(r):
    return (str(r['date']), (r['account_type'] or '').strip(), (r['category'] or '').strip(),
            (r['details'] or '').strip(), int(round(r['income'] or 0)),
            int(round(r['expense'] or 0)))


def _db_tr_key(r):
    return (str(r['date']), (r['from_account'] or '').strip(),
            (r['to_account'] or '').strip(), int(round(r['amount'] or 0)))


def examine(db, plan):
    """すでに入っているものと突き合わせて、重複の数と、新しい科目・口座を数える。

    まったく同じ内容の取引が2件ある帳面があるので、「いくつあるか」で数える。
    すでに1件あって Excel に2件あるなら、足りない1件だけを入れる。
    """
    from collections import Counter
    have_tx, have_tr = Counter(), Counter()
    for y in sorted(plan.fiscal_years):
        for r in db.get_transactions(y):
            have_tx[_db_tx_key(r)] += 1
        for r in db.get_transfers(y):
            have_tr[_db_tr_key(r)] += 1

    plan.dup_transactions = 0
    for t in plan.transactions:
        k = _tx_key(t)
        if have_tx[k] > 0:
            have_tx[k] -= 1
            t['dup'] = True
            plan.dup_transactions += 1
        else:
            t['dup'] = False
    plan.dup_transfers = 0
    for t in plan.transfers:
        k = _tr_key(t)
        if have_tr[k] > 0:
            have_tr[k] -= 1
            t['dup'] = True
            plan.dup_transfers += 1
        else:
            t['dup'] = False

    cats = set(db.get_category_names())
    accs = set(db.get_account_names())
    new_cats, new_accs = [], []
    for t in plan.transactions:
        if t.get('dup'):
            continue
        c = (t.get('category') or '').strip()
        if c and c not in cats and c not in new_cats:
            new_cats.append(c)
        a = (t.get('account') or '').strip()
        if a and a not in accs and a not in new_accs:
            new_accs.append(a)
    for t in plan.transfers:
        if t.get('dup'):
            continue
        for a in ((t.get('from_account') or '').strip(), (t.get('to_account') or '').strip()):
            if a and a not in accs and a not in new_accs:
                new_accs.append(a)
    for a in plan.openings:
        if a and a not in accs and a not in new_accs:
            new_accs.append(a)
    plan.new_categories = new_cats
    plan.new_accounts = new_accs
    return plan


# ============================================================
# 入れる
# ============================================================
def _columns(conn, table):
    return {r[1] for r in conn.execute(f'PRAGMA table_info({table})').fetchall()}


def apply_plan(db, plan, take_openings=True):
    """まとめて入れる。途中で失敗したらぜんぶ取り消す。戻り: 入れた件数の内訳"""
    conn = db.conn
    tx_cols = _columns(conn, 'transactions')
    tr_cols = _columns(conn, 'transfers')
    has_event = 'event' in tx_cols                     # 新しい版には「行事」の列がある
    added_tx = added_tr = 0
    try:
        conn.execute('BEGIN')
        for name in plan.new_categories:
            conn.execute('INSERT OR IGNORE INTO categories (name,sort_order)'
                         ' VALUES (?,(SELECT COALESCE(MAX(sort_order)+1,0) FROM categories))',
                         (name,))
        for name in plan.new_accounts:
            conn.execute('INSERT OR IGNORE INTO accounts (name,sort_order)'
                         ' VALUES (?,(SELECT COALESCE(MAX(sort_order)+1,0) FROM accounts))',
                         (name,))
        for y in sorted(plan.fiscal_years):
            conn.execute('INSERT OR IGNORE INTO fiscal_years (year) VALUES (?)', (y,))

        for t in plan.transactions:
            if t.get('dup'):
                continue
            cols = ['fiscal_year', 'date', 'account_type', 'category', 'details',
                    'income', 'expense', 'memo']
            vals = [fiscal_year_of(t['date']), t['date'], t['account'], t['category'],
                    t['details'], t['income'], t['expense'], t['memo']]
            if has_event:
                cols.append('event')
                vals.append(t.get('event', ''))
            conn.execute(f"INSERT INTO transactions ({','.join(cols)})"
                         f" VALUES ({','.join('?' * len(cols))})", vals)
            added_tx += 1

        for t in plan.transfers:
            if t.get('dup'):
                continue
            cols = ['fiscal_year', 'date', 'from_account', 'to_account', 'amount', 'details']
            vals = [fiscal_year_of(t['date']), t['date'], t['from_account'], t['to_account'],
                    t['amount'], t['details']]
            if 'memo' in tr_cols:
                cols.append('memo')
                vals.append(t.get('memo', ''))
            conn.execute(f"INSERT INTO transfers ({','.join(cols)})"
                         f" VALUES ({','.join('?' * len(cols))})", vals)
            added_tr += 1

        if take_openings and plan.openings:
            for y in sorted(plan.fiscal_years):
                for name, amount in plan.openings.items():
                    conn.execute(
                        'INSERT INTO account_balances (fiscal_year,account_name,balance)'
                        ' VALUES (?,?,?)'
                        ' ON CONFLICT(fiscal_year,account_name)'
                        ' DO UPDATE SET balance=excluded.balance',
                        (y, name, amount))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {'transactions': added_tx, 'transfers': added_tr,
            'dup_transactions': plan.dup_transactions, 'dup_transfers': plan.dup_transfers,
            'new_categories': plan.new_categories, 'new_accounts': plan.new_accounts,
            'openings': len(plan.openings) if take_openings else 0}


# ============================================================
# 画面（tkinter）
# ============================================================
def import_excel_dialog(app, path=None):
    """「データ ▼ → Excelを取り込み」から呼ぶ。app は .root / .db / .current_fy を持つもの"""
    from tkinter import filedialog, messagebox

    root = getattr(app, 'root', None)
    db = app.db
    if not EXCEL_AVAILABLE:
        messagebox.showerror('Excel取込',
                             'openpyxl が入っていないため、Excel を読み込めません。',
                             parent=root)
        return

    if path is None:
        path = filedialog.askopenfilename(
            title='取り込む Excel ファイルを選択',
            filetypes=[('Excel ブック', '*.xlsx *.xlsm'), ('すべてのファイル', '*.*')],
            parent=root)
    if not path:
        return

    try:
        fy = app.current_fy.get() if hasattr(app, 'current_fy') else None
    except Exception:
        fy = None

    try:
        plan = read_workbook(path, fy)
    except Exception as e:
        messagebox.showerror('Excel取込',
                             'ファイルを読み込めませんでした。\n\n'
                             '.xlsx か .xlsm を選んでください。\n'
                             '（.xls のときは Excel で開いて .xlsx で保存しなおしてください）\n\n'
                             f'{e}', parent=root)
        return

    if not plan.transactions and not plan.transfers:
        messagebox.showwarning('Excel取込',
                               '取り込める行が見つかりませんでした。\n\n'
                               '「日付」と「収入・支出（または金額）」の見出しがある表を'
                               '入れてください。', parent=root)
        return

    examine(db, plan)

    years = '・'.join(f'{y}年度' for y in sorted(plan.fiscal_years)) or '（年度不明）'
    msg = [f'{years} のデータを取り込みます。', '']
    msg.append(f'取引 {plan.add_transactions} 件、振替 {plan.add_transfers} 件 を追加します。')
    if plan.dup_transactions or plan.dup_transfers:
        msg.append(f'すでに入っている 取引 {plan.dup_transactions} 件、'
                   f'振替 {plan.dup_transfers} 件 はスキップします。')
    if plan.skipped:
        msg.append(f'{len(plan.skipped)} 行は形式不正のためスキップします'
                   '（表題や合計の行なら、そのままで大丈夫です）。')
    if plan.new_categories:
        msg.append('新しい科目を追加: ' + '、'.join(plan.new_categories))
    if plan.new_accounts:
        msg.append('新しい口座を追加: ' + '、'.join(plan.new_accounts))
    if plan.openings:
        msg.append('期首残高も取り込みます: '
                   + '、'.join(f'{k} {v:,}' for k, v in plan.openings.items()))
    msg += ['', '念のため取込前のバックアップを推奨します。', '', '続けますか？']

    if not messagebox.askyesno('Excel取込の確認', '\n'.join(msg), parent=root):
        return

    try:
        res = apply_plan(db, plan)
    except Exception as e:
        messagebox.showerror('Excel取込',
                             f'取込中にエラーが発生したため、すべて取り消しました。\n\n{e}',
                             parent=root)
        return

    for name in ('refresh_list', 'refresh_summary', 'refresh_balance',
                 'refresh_category_combos', 'refresh_account_combos', 'refresh_compare_list'):
        fn = getattr(app, name, None)
        if callable(fn):
            try:
                fn()
            except Exception:
                pass

    done = [f"取引 {res['transactions']} 件、振替 {res['transfers']} 件 を取り込みました。"]
    if res['dup_transactions'] or res['dup_transfers']:
        done.append(f"重複スキップ: 取引 {res['dup_transactions']} 件、"
                    f"振替 {res['dup_transfers']} 件")
    if plan.skipped:
        done.append(f'形式不正スキップ: {len(plan.skipped)} 行')
    if res['new_categories']:
        done.append('追加した科目: ' + '、'.join(res['new_categories']))
    if res['new_accounts']:
        done.append('追加した口座: ' + '、'.join(res['new_accounts']))
    messagebox.showinfo('Excel取込完了', '\n'.join(done), parent=root)
